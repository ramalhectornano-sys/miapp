from io import BytesIO
import zipfile
import httpx
from fastapi import HTTPException

from ..db import Database, User


async def download_resume(db: Database, user: User, resume: dict, service=False) -> bytes:
    path = resume.get("storage_path") or resume.get("object_path")
    if not path or not path.startswith(f"{user.id}/") or ".." in path.split("/"):
        raise HTTPException(409, "La ruta del currículum no es válida.")
    from urllib.parse import quote
    try:
        async with db.client.stream("GET", f"{db.settings.supabase_url}/storage/v1/object/authenticated/resumes/{quote(path, safe='/')}",
                headers=db.headers(None if service else user)) as response:
            if response.status_code != 200:
                raise HTTPException(502, "No se pudo descargar el currículum privado.")
            chunks, length = [], 0
            async for chunk in response.aiter_bytes():
                length += len(chunk)
                if length > db.settings.max_resume_bytes:
                    raise HTTPException(413, "El currículum supera el máximo de 10 MB.")
                chunks.append(chunk)
            return b"".join(chunks)
    except httpx.RequestError:
        raise HTTPException(502, "No se pudo descargar el currículum.")


def extract_text(content: bytes, filename: str) -> str:
    try:
        if filename.lower().endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise HTTPException(422, "El PDF tiene contraseña. Cargá una copia sin contraseña.")
            if len(reader.pages) > 100:
                raise HTTPException(422, "El currículum no puede superar las 100 páginas.")
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif filename.lower().endswith(".docx"):
            from docx import Document
            with zipfile.ZipFile(BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024 or len(archive.infolist()) > 500:
                    raise HTTPException(422, "El documento comprimido excede los límites de procesamiento.")
            doc = Document(BytesIO(content))
            text = "\n".join([p.text for p in doc.paragraphs] + [" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows])
        else:
            raise HTTPException(422, "Usá un currículum PDF o DOCX.")
        return text.strip()[:50000]
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "No se pudo leer el documento. Verificá el archivo.")


async def extract_resume(db: Database, user: User, resume_id: str):
    resume = await db.owned("resume_versions", resume_id, user)
    data = await download_resume(db, user, resume)
    import asyncio
    text = await asyncio.to_thread(extract_text, data, resume.get("original_filename") or resume.get("file_name") or resume.get("filename") or resume.get("object_path", ""))
    return {"resume_id": resume_id, "text": text, "extraction_method": "document_text", "profile_saved": False,
        "requires_confirmation": True, "message": "Se extrajo texto del archivo. Revisá los datos antes de incorporarlos al perfil; no se agregó información inferida." if text else "No se detectó texto. Los PDF escaneados necesitan OCR; completá el perfil manualmente."}
